"""Generate the synthetic inspection export used by the demo.

Everything here is invented: the property, the units, the residents, the
inspectors, the addresses and the work order numbers. The rough-note *style*
is modelled on how maintenance technicians actually write, because that is
the problem the tool exists to solve, but no real inspection data is used.

The inspection photographs are drawn from scratch too — flat shapes and a
caption, stamped SAMPLE - SYNTHETIC DATA. They exist so the photo half of the
pipeline (extraction from xl/media, anchor matching, resampling, signature
filtering) actually runs in the public sample. No real photograph is included.

Run:  python sample-data/generate_sample.py
"""
from __future__ import annotations

import random
import shutil
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image, ImageDraw, ImageFont

SEED = 4820
OUT = Path(__file__).parent / "Prop 4820 - Inspection Report.xlsx"

PROPERTY_CODE = "4820"
ADDRESS = "1400 Alderpath Ave. #04 Torrance, CA - 90503 us"

# The fixed annual checklist. Wording matches the export format, typos included.
ITEMS = [
    "Front door has working locks, unit # listed and weatherstripping is in tact",
    "Smoke alarms and CO2 alarms are in working condition and placed in required areas of unit",
    "Stove/ Refrigerator/ Dishwasher/ Microwave are in working condition",
    "Kitchen cabinets and countertops are in good working condition",
    "All plumbing including water heater,toilets, faucets, under sinks, shower heads are free from leaks",
    "Windows and screens are in tact and in working condition",
    "Paint, ceiling and walls are in good condition and free from water damage and signs of leaks",
    "Carpet and flooring is in good condition and free from trip hazards",
    "Wall heater and Water heater is in working condition and free from furniture or belongings blocking it",
]

# Rough field notes: dropped articles, misspellings, two issues in one note,
# and several that state a fix instead of a condition. The narrative step has
# to turn each of these into one clean factual sentence.
NOTES = {
    0: ["Needs new weather strips", "Dead bolt sticks hard to lock", "Unit number missing from door"],
    1: ["Replace all missing smoke detectors", "Missing smoke detectors in hallway", "Co2 alarm chirping needs battery"],
    2: ["Stove wont turn on right side", "Fridge door seal torn", "Dishwaser not draining"],
    3: ["Kitchen drawer needs repair", "Fix hinges from kitchen cabinets", "Countertop paint coming off"],
    4: ["Tub sprout needs caulking in hallway bathroom", "Shower handle broken", "Bathroom sink drains slow",
        "Garbage disposal not working Bathtub and bathroom sink need strainers", "Water heater rusted at bottom"],
    5: ["Replace window screen in kitchen", "Old blinds were left behind", "Bedroom window wont stay up"],
    6: ["Drywall repair in dining room area behind kitchen wall",
        "Ceiling fan and dining room is broken and not secured", "Water stain on ceiling by bathroom"],
    7: ["Carpet in bedrooms are old and damaged. Need to be replaced with hardwood floor.",
        "Floor planks are lifting in living room", "Hallway floor damage", "Floor tack strip is missing screws"],
    8: ["Living room Wall heater knob broken", "Water heater closet blocked with boxes"],
}
ACTIONS = {0: "Repair", 1: "Replace", 2: "Repair", 3: "Repair", 4: "Repair",
           5: "Repair", 6: "Repair", 7: "Repair", 8: "Repair"}

INSPECTORS = ["Marisol Vega", "Daniel Okonkwo", "Priya Raman", "Tomas Herrera"]
RESIDENTS = ["Avery Lindqvist", "Marcus Bello", "Sofia Restrepo", "Nathan Oyelaran",
             "Delia Marchetti", "Ruben Castellanos", "Harriet Nwosu", "Jonas Petrauskas",
             "Camille Oduya", "Theo Vasquez", "", "", ""]
NO_ACCESS_REASONS = ["No key on file, request from PM", "Refused access, dog in unit",
                     "Vacant, unit under turn", "Tenant not home, second attempt"]

COL_HEADERS = [(1, " "), (3, " Results"), (4, " Observation"), (6, " Work Order"),
               (7, " Charge"), (10, " Responsibility"), (11, " Detail Level Notes"), (13, " Photo")]


# ---------------------------------------------------------------------------
# Synthetic photographs
#
# Yardi embeds inspection photos in the workbook itself: the image binary lives
# in xl/media and an anchor in xl/drawings positions it on the row of the
# checklist item it belongs to, in the Photo column (M). A signature pad is
# embedded the same way but lands in a low column near the Signatures label.
#
# The sample therefore has to embed real images, or the photo half of the
# pipeline never runs. These are drawn from scratch: flat shapes, a caption and
# a "SAMPLE - SYNTHETIC DATA" stamp. They are not photographs of anything and
# they are not derived from any real inspection.
#
# Two details matter for the parser:
#   * every image must be byte-unique, because the parser drops duplicates by
#     md5 hash (a real Yardi quirk it has to defend against), and
#   * the anchor column must be M or later, or the photo is read as a signature.
# ---------------------------------------------------------------------------

FONT_DIR = Path(__file__).parent.parent / "assets" / "fonts"
STAMP = "SAMPLE - SYNTHETIC DATA"

# Three source sizes, so the resampling code is exercised in all three of its
# branches: upscale-with-sharpen (the common Yardi thumbnail case), pass
# through, and downscale.
SIZE_THUMB = (260, 195)     # below TARGET_DISPLAY_WIDTH -> upscale path
SIZE_MID = (960, 720)       # inside the band -> pass through
SIZE_LARGE = (2400, 1800)   # above MAX_WIDTH -> downscale path

SIG_SIZE = (480, 110)       # 4.36:1, wide enough to trip the aspect-ratio test

# Muted, deliberately flat interior tones. Nothing here is meant to look like a
# photograph; it is meant to occupy the photo slot honestly.
WALL_TONES = [(214, 209, 198), (203, 205, 199), (219, 212, 202), (198, 201, 205)]
FLOOR_TONES = [(150, 132, 110), (128, 124, 118), (163, 148, 126), (112, 110, 106)]
MARK = (194, 112, 61)  # the copper accent, used to ring the area of interest


def _font(size: int, bold: bool = False):
    name = "Ubuntu-Bold.ttf" if bold else "Ubuntu-Regular.ttf"
    try:
        return ImageFont.truetype(str(FONT_DIR / name), size)
    except Exception:
        return ImageFont.load_default()


# Short area names for the photo caption, one per checklist item, in ITEMS
# order. Truncating the checklist question itself reads as a cut-off sentence.
ITEM_LABELS = [
    "Front door", "Smoke & CO alarms", "Appliances", "Kitchen", "Plumbing",
    "Windows", "Paint & walls", "Flooring", "Heating",
]


def _make_unit_photo(out_path: Path, unit: str, label: str, serial: int, size) -> None:
    rng = random.Random(f"{unit}-{serial}")
    w, h = size
    img = Image.new("RGB", (w, h), rng.choice(WALL_TONES))
    d = ImageDraw.Draw(img)

    # Floor band along the bottom, and a simple rectangle standing in for the
    # fixture or surface the note is about.
    floor_y = int(h * rng.uniform(0.62, 0.74))
    d.rectangle([0, floor_y, w, h], fill=rng.choice(FLOOR_TONES))

    bw, bh = int(w * rng.uniform(0.22, 0.40)), int(h * rng.uniform(0.26, 0.46))
    bx = int(w * rng.uniform(0.08, 0.55))
    by = floor_y - bh
    shade = rng.randint(150, 205)
    d.rectangle([bx, by, bx + bw, by + bh], fill=(shade, shade - 6, shade - 14),
                outline=(90, 88, 84), width=max(1, w // 400))

    # The area of interest, ringed in the report's accent colour.
    rx = bx + int(bw * rng.uniform(0.15, 0.6))
    ry = by + int(bh * rng.uniform(0.15, 0.6))
    rr = int(min(w, h) * rng.uniform(0.07, 0.13))
    d.ellipse([rx - rr, ry - rr, rx + rr, ry + rr], outline=MARK, width=max(2, w // 260))

    # Per-image speckle. Cheap, and it guarantees no two images hash alike.
    for _ in range(24):
        px, py = rng.randrange(w), rng.randrange(h)
        d.point((px, py), fill=(rng.randrange(256), rng.randrange(256), rng.randrange(256)))

    # Caption bar.
    bar_h = max(22, int(h * 0.14))
    d.rectangle([0, h - bar_h, w, h], fill=(31, 78, 70))
    f_small = _font(max(10, int(bar_h * 0.34)))
    f_stamp = _font(max(9, int(bar_h * 0.28)), bold=True)
    d.text((int(w * 0.02), h - bar_h + int(bar_h * 0.10)),
           f"UNIT {unit} - {label}", fill=(245, 245, 240), font=f_small)
    d.text((int(w * 0.02), h - bar_h + int(bar_h * 0.55)),
           STAMP, fill=(194, 112, 61), font=f_stamp)

    img.save(out_path, "JPEG", quality=82, optimize=True)


def _make_signature(out_path: Path, inspector: str, serial: int) -> None:
    """A scribble in a wide box. The pipeline should classify this as a
    signature and keep it out of the owner-facing photo grids."""
    rng = random.Random(f"sig-{inspector}-{serial}")
    w, h = SIG_SIZE
    img = Image.new("RGB", (w, h), (252, 252, 250))
    d = ImageDraw.Draw(img)
    x, y = int(w * 0.06), int(h * 0.62)
    pts = [(x, y)]
    for i in range(14):
        x += int(w * 0.062)
        y = int(h * (0.62 + rng.uniform(-0.34, 0.20)))
        pts.append((x, y))
    d.line(pts, fill=(28, 42, 84), width=3, joint="curve")
    d.line([int(w * 0.05), int(h * 0.86), int(w * 0.95), int(h * 0.86)], fill=(170, 170, 165), width=1)
    d.text((int(w * 0.05), int(h * 0.02)), STAMP, fill=(190, 190, 185), font=_font(11, bold=True))
    img.save(out_path, "JPEG", quality=80, optimize=True)


def _photo_size(serial: int):
    if serial % 13 == 0:
        return SIZE_LARGE
    if serial % 3 == 0:
        return SIZE_MID
    return SIZE_THUMB


def build() -> None:
    random.seed(SEED)
    # A separate stream for photo decisions, so adding photos does not shift
    # the main sequence and change which checklist items fail.
    photo_rng = random.Random(SEED + 1)
    tmp_dir = Path(tempfile.mkdtemp(prefix="sample-photos-"))
    serial = 0

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    row = 1
    inspection_id = 51200
    work_order = 412600
    units = [f"{n:02d}" for n in range(1, 29)]
    # four units cannot be entered
    no_access = {units[3], units[11], units[19], units[25]}

    for i, unit in enumerate(units):
        inspector = INSPECTORS[i % len(INSPECTORS)]
        day = 11 + (i % 5)
        hour = 9 + (i % 9)
        minute = (i * 7) % 60
        stamp = f"{inspector} 5/{day}/2026 {(hour % 12) or 12}:{minute:02d} {'PM' if hour >= 12 else 'AM'}"

        ws.cell(row, 1, "Inspection ID"); ws.cell(row, 3, inspection_id); row += 1
        ws.cell(row, 1, "Inspector"); ws.cell(row, 3, inspector); row += 1
        ws.cell(row, 1, "Inspected Date"); ws.cell(row, 3, f"5/{day}/2026 {hour}:{minute:02d}"); row += 1
        ws.cell(row, 1, "Entity Type"); ws.cell(row, 3, "Unit"); row += 1
        ws.cell(row, 1, "Tenant Name"); ws.cell(row, 3, RESIDENTS[i % len(RESIDENTS)]); row += 1
        # Yardi merges these two labels into one cell, newline separated
        ws.cell(row, 1, "Property Code\nUnit Code")
        ws.cell(row, 3, f"{PROPERTY_CODE}    \n{unit}  "); row += 1
        ws.cell(row, 1, "Unit Address"); ws.cell(row, 3, ADDRESS); row += 1
        for col, val in COL_HEADERS:
            ws.cell(row, col, val)
        row += 1
        ws.cell(row, 1, "Annual Maintenace Unit Inspections"); row += 1

        if unit in no_access:
            reason = random.choice(NO_ACCESS_REASONS)
            ws.cell(row, 1, ITEMS[0]); ws.cell(row, 3, "Cancel")
            ws.cell(row, 11, f"[{stamp}] {reason}"); row += 1
            overall = "Cancel - Unable to inspect"
        else:
            # every checklist item is answered; failures carry an action and a note
            n_fail = random.choice([0, 0, 1, 1, 1, 2, 2, 3, 5])
            failed = set(random.sample(range(len(ITEMS)), n_fail)) if n_fail else set()
            for idx, question in enumerate(ITEMS):
                ws.cell(row, 1, question)
                if idx in failed:
                    ws.cell(row, 3, "No")
                    ws.cell(row, 4, ACTIONS[idx])
                    ws.cell(row, 6, work_order)
                    ws.cell(row, 11, f"[{stamp}] {random.choice(NOTES[idx])}")
                    ws.cell(row, 13, " ")
                    # Photos anchor on the item's own row, in the Photo column.
                    for _ in range(photo_rng.choice([1, 1, 1, 2, 2, 3])):
                        serial += 1
                        img_path = tmp_dir / f"p{serial:04d}.jpg"
                        _make_unit_photo(img_path, unit, ITEM_LABELS[idx], serial, _photo_size(serial))
                        ws.add_image(XLImage(str(img_path)), f"M{row}")
                else:
                    ws.cell(row, 3, "Yes")
                row += 1
            overall = "Pass - Corrective action needed" if failed else "Pass - No action needed"

        ws.cell(row, 1, "Overall Result"); ws.cell(row, 3, overall); row += 1
        ws.cell(row, 1, "Overall Notes")
        if unit not in no_access and not failed:
            # A clean unit still gets walked and photographed. These anchor off
            # any checklist row, so they arrive as supplemental photos.
            for _ in range(2):
                serial += 1
                img_path = tmp_dir / f"p{serial:04d}.jpg"
                _make_unit_photo(img_path, unit, "General walkthrough", serial, _photo_size(serial))
                ws.add_image(XLImage(str(img_path)), f"M{row}")
        row += 1
        ws.cell(row, 1, "Tenant Response"); row += 1
        ws.cell(row, 2, "Signatures"); ws.cell(row, 3, "Inspector Signature ")
        ws.cell(row, 6, "Owner Signature"); ws.cell(row, 8, "Tenant Signature"); row += 1
        ws.cell(row, 3, " ")
        if unit not in no_access:
            serial += 1
            sig_path = tmp_dir / f"s{serial:04d}.jpg"
            _make_signature(sig_path, inspector, serial)
            ws.add_image(XLImage(str(sig_path)), f"C{row}")
        row += 2

        inspection_id += 1
        work_order += 1

    wb.save(OUT)
    shutil.rmtree(tmp_dir, ignore_errors=True)
    size_kb = OUT.stat().st_size / 1024
    print(f"wrote {OUT}  ({len(units)} units, {len(no_access)} not inspected, "
          f"{serial} embedded images, {size_kb:.0f} KB)")


if __name__ == "__main__":
    build()
